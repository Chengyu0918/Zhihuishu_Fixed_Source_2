#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
读取三个路径Excel表格，对比数据库中节点的层级信息
第一列：出发节点
第二列：目的节点
第三列：出发节点的层级（从数据库查询）
第四列：目的节点的层级（从数据库查询）
"""

import sqlite3
import pandas as pd
from openpyxl import load_workbook

# 数据库路径
DATABASE_PATH = 'database.db'

# 三个Excel文件路径（输入和输出）
EXCEL_FILES = [
    ('data/2/电子信息工程主路径梳理（v1）.xlsx', 'data/2/电子信息工程主路径梳理（v1）_层级.xlsx'),
    ('data/2/电子信息工程兴趣路径梳理（v1）.xlsx', 'data/2/电子信息工程兴趣路径梳理（v1）_层级.xlsx'),
    ('data/2/电子信息工程拓展路径梳理（v1）.xlsx', 'data/2/电子信息工程拓展路径梳理（v1）_层级.xlsx')
]

# 层级映射（字符串转数字）
LAYER_TO_NUM = {
    'LAYER_ABILITY': 1,
    'LAYER_PROBLEM': 2,
    'LAYER_PROFESSIONAL': 3,
    'LAYER_DISCIPLINE_BASE': 4,
    'LAYER_ADVANCED_BASE': 5,
    'LAYER_MATH_PHYSICS': 6
}


def get_node_layer(cursor, node_name):
    """从数据库查询节点的层级"""
    cursor.execute("SELECT layer FROM nodes WHERE name = ?", (node_name,))
    result = cursor.fetchone()
    if result:
        layer_str = result[0]
        # 转换为数字
        return LAYER_TO_NUM.get(layer_str, layer_str)
    return None


def process_excel_file(input_path, output_path, cursor):
    """处理单个Excel文件的所有Sheet"""
    print(f"\n处理文件: {input_path}")
    print(f"  输出到: {output_path}")
    
    # 加载工作簿
    wb = load_workbook(input_path)
    
    for sheet_name in wb.sheetnames:
        print(f"  处理Sheet: {sheet_name}")
        ws = wb[sheet_name]
        
        # 获取最大行数
        max_row = ws.max_row
        
        # 遍历每一行（从第2行开始，假设第1行是标题）
        for row in range(2, max_row + 1):
            # 获取第一列和第二列的值
            col1_value = ws.cell(row=row, column=1).value
            col2_value = ws.cell(row=row, column=2).value
            
            if col1_value:
                # 查询第一列节点的层级
                layer1 = get_node_layer(cursor, str(col1_value).strip())
                ws.cell(row=row, column=3).value = layer1
            
            if col2_value:
                # 查询第二列节点的层级
                layer2 = get_node_layer(cursor, str(col2_value).strip())
                ws.cell(row=row, column=4).value = layer2
        
        print(f"    已处理 {max_row - 1} 行数据")
    
    # 保存到新文件
    wb.save(output_path)
    print(f"  文件已保存: {output_path}")


def main():
    """主函数"""
    print("=" * 60)
    print("读取Excel表格并写入节点层级信息")
    print("=" * 60)
    
    # 连接数据库
    conn = sqlite3.connect(DATABASE_PATH)
    cursor = conn.cursor()
    
    import os
    
    # 处理每个Excel文件
    for input_path, output_path in EXCEL_FILES:
        try:
            if os.path.exists(input_path):
                print(f"文件存在: {input_path}")
            else:
                print(f"文件不存在: {input_path}")
                continue
            process_excel_file(input_path, output_path, cursor)
        except Exception as e:
            import traceback
            print(f"  处理文件出错: {e}")
            traceback.print_exc()
    
    conn.close()
    
    print("\n" + "=" * 60)
    print("处理完成！")
    print("=" * 60)


if __name__ == '__main__':
    main()
